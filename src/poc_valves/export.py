from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

from .models import (
    DocumentResult,
    FieldDefinition,
    SchemaDefinition,
)


class ExportError(RuntimeError):
    """Error raised during export operations."""
    pass


def write_results_json(
    output_path: str | Path, result: DocumentResult
) -> Path:
    """Write extraction results to a JSON file."""
    output_path = Path(output_path)
    output_path.write_text(
        json.dumps(
            result.to_dict(), indent=2, ensure_ascii=True
        ),
        encoding="utf-8",
    )
    return output_path


def write_results_csv(
    output_path: str | Path,
    schema: SchemaDefinition,
    result: DocumentResult,
) -> Path:
    """Write extraction results to a CSV file."""
    output_path = Path(output_path)
    rows: list[dict[str, object]] = []
    field_index = {
        field.field_id: field for field in schema.fields
    }
    headers = [
        "field_id",
        "field_name",
        "subcategory",
        "value",
        "normalized_value",
        "source_page",
        "confidence",
        "notes",
    ]
    for field_id, extracted in result.fields.items():
        field_def = field_index.get(field_id)
        rows.append(
            {
                "field_id": field_id,
                "field_name": (
                    field_def.field_name if field_def else ""
                ),
                "subcategory": (
                    field_def.subcategory if field_def else ""
                ),
                "value": extracted.value,
                "normalized_value": extracted.normalized_value,
                "source_page": extracted.source_page,
                "confidence": extracted.confidence,
                "notes": extracted.notes,
            }
        )

    with output_path.open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        if rows:
            writer.writerows(rows)
    return output_path


def write_results_xlsx(
    output_path: str | Path,
    schema: SchemaDefinition,
    result: DocumentResult,
) -> Path:
    """Write extraction results to an Excel file."""
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover - optional dependency
        raise ExportError(
            "openpyxl is required for Excel export. "
            "Install the 'excel' extra."
        ) from exc

    output_path = Path(output_path)
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    grouped_fields: dict[str, list[FieldDefinition]] = (
        defaultdict(list)
    )
    for field in schema.fields:
        grouped_fields[field.subcategory].append(field)

    for subcategory, fields in sorted(
        grouped_fields.items()
    ):
        sheet = workbook.create_sheet(
            title=subcategory[:31] or "Sheet"
        )
        sheet.append(
            [
                "field_id",
                "field_name",
                "value",
                "normalized_value",
                "source_page",
                "confidence",
                "notes",
            ]
        )
        for field in fields:
            extracted = result.fields.get(field.field_id)
            sheet.append(
                [
                    field.field_id,
                    field.field_name,
                    (
                        None
                        if extracted is None
                        else extracted.value
                    ),
                    (
                        None
                        if extracted is None
                        else extracted.normalized_value
                    ),
                    (
                        None
                        if extracted is None
                        else extracted.source_page
                    ),
                    (
                        None
                        if extracted is None
                        else extracted.confidence
                    ),
                    (
                        None
                        if extracted is None
                        else extracted.notes
                    ),
                ]
            )

    meta = workbook.create_sheet(title="metadata")
    meta.append(["key", "value"])
    meta.append(["document_name", result.document_name])
    for key, value in result.metadata.items():
        meta.append([key, value])
    meta.append(
        ["validation_issues", len(result.validation)]
    )

    workbook.save(output_path)
    return output_path