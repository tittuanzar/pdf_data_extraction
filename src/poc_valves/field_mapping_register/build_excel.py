from __future__ import annotations

from pathlib import Path
from typing import Any

from .models import ReviewItem


LONG_TEXT_COLUMNS = {
    "mapping_logic",
    "notes_and_additional_information",
    "validations",
}


def _column_width(value: Any) -> int:
    """Calculate the column width for a value."""
    if value is None:
        return 12
    text = str(value)
    if "\n" in text:
        return min(
            max(len(line) for line in text.splitlines()) + 2,
            60,
        )
    return min(max(len(text) + 2, 12), 42)


def _write_header(
    sheet, headers: list[str]
) -> None:
    """Write and format the header row."""
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            wrap_text=True, vertical="top"
        )


def build_excel(
    output_path: str | Path,
    rows: list[dict[str, Any]],
    *,
    review_items: list[ReviewItem] | None = None,
) -> Path:
    """Build an Excel file from extracted rows."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except Exception as exc:  # pragma: no cover - optional dependency
        raise RuntimeError(
            "openpyxl is required to write Excel output. "
            "Install the 'excel' extra."
        ) from exc

    output_path = Path(output_path)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "field_mapping_register"

    if rows:
        headers = list(rows[0].keys())
    else:
        headers = []
    _write_header(ws, headers)

    for row in rows:
        ws.append(
            [row.get(header) for header in headers]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for index, header in enumerate(headers, start=1):
        max_width = _column_width(header)
        for row in rows[:200]:
            max_width = max(
                max_width, _column_width(row.get(header))
            )
        if header in LONG_TEXT_COLUMNS:
            for cell in ws.iter_rows(
                min_col=index,
                max_col=index,
                min_row=2,
            ):
                cell[0].alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
        ws.column_dimensions[
            ws.cell(row=1, column=index).column_letter
        ].width = min(max_width, 60)

    if review_items:
        review = workbook.create_sheet("review")
        review_headers = [
            "serial_number",
            "reason",
            "row_json",
        ]
        _write_header(review, review_headers)
        for item in review_items:
            review.append(
                [
                    item.serial_number,
                    item.reason,
                    str(item.row),
                ]
            )
        review.freeze_panes = "A2"
        review.auto_filter.ref = review.dimensions
        for col in ("A", "B", "C"):
            review.column_dimensions[col].width = (
                40 if col == "C" else 24
            )

    workbook.save(output_path)
    return output_path